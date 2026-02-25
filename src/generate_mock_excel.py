import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from config import MASTER_EXCEL_PATH, VALIDATION_RULES

def create_mock_excel():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    
    header_fill = PatternFill(start_color="800000", end_color="800000", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    def setup_sheet(sheet_name, headers, rows):
        ws = wb.create_sheet(title=sheet_name)
        # Add headers to row 4
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
            
        # Add data starting at row 5
        for r_idx, row in enumerate(rows, 5):
            for c_idx, val in enumerate(row, 1):
                ws.cell(row=r_idx, column=c_idx, value=val)
                
    # 01. Scorecards
    setup_sheet("DIV_01_SCORECARDS", 
                VALIDATION_RULES["DIV_01_SCORECARDS"]["required_columns"],
                [
                    ["2026-01", 1250, 450.5, 120.0],
                    ["2026-02", 1420, 510.2, 145.5],
                    ["2026-03", 1600, 580.0, 160.0]
                ])

    # 02. Districts
    districts_data = []
    districts = VALIDATION_RULES["DIV_02_DISTRICTS"]["district_names"]
    for i, d in enumerate(districts):
        # Generate varied recruitment numbers so charts look organic
        jan = 200 + (i * 15) if i % 2 == 0 else 500 - (i * 10)
        feb = jan + 20 + (i * 5)
        
        # Manually spike Nuwara Eliya (index 19) to match user document requirements
        if d == "Nuwara Eliya":
            jan, feb = 4500, 5937
            
        tot = jan + feb
        districts_data.append([d, jan, feb, 0,0,0,0,0,0,0,0,0,0, tot, i+1])
    setup_sheet("DIV_02_DISTRICTS", VALIDATION_RULES["DIV_02_DISTRICTS"]["required_columns"], districts_data)

    # 03. Financials
    setup_sheet("DIV_03_FINANCIALS", VALIDATION_RULES["DIV_03_FINANCIALS"]["required_columns"], [
        ["වාරික ආදායම (Premium Income)", "Income", 450.5, 510.2, 0,0,0,0,0,0,0,0,0,0, 960.7],
        ["ආයෝජන ආදායම (Investment Income)", "Income", 120.0, 125.5, 0,0,0,0,0,0,0,0,0,0, 245.5],
        ["බැංකු හරහා එකතු කිරීම් (Bank Collections)", "Income", 300.0, 350.0, 0,0,0,0,0,0,0,0,0,0, 650.0],
        ["තැපැල් කාර්යාල හරහා (Post Office)", "Income", 150.0, 160.0, 0,0,0,0,0,0,0,0,0,0, 310.0],
        ["පරිපාලන වියදම් (Admin Expenses)", "Expense", 45.0, 48.0, 0,0,0,0,0,0,0,0,0,0, 93.0],
        ["විශ්‍රාම වැටුප් ගෙවීම් (Pension Payouts)", "Expense", 200.0, 210.0, 0,0,0,0,0,0,0,0,0,0, 410.0],
        ["ප්‍රාග්ධන වියදම් (Capital Expenditure)", "Expense", 15.0, 5.0, 0,0,0,0,0,0,0,0,0,0, 20.0]
    ])

    # 04. Board
    setup_sheet("DIV_04_BOARD", VALIDATION_RULES["DIV_04_BOARD"]["required_columns"], [
        ["සභාපති", "Chairman", "සභාපති (Chairman)", "දශක දෙකකට වැඩි කාලයක් රාජ්‍ය සේවයේ සහ මූල්‍ය පරිපාලනයේ අත්දැකීම් සහිතයි.", "DIRECTOR_01.jpg", 1],
        ["අධ්‍යක්ෂ ජනරාල්", "Director General", "අධ්‍යක්ෂ ජනරාල් (Director Gen)", "කළමනාකරණ පරිපාලනය පිළිබඳ විශේෂඥතාවයක් ලබා ඇත.", "DIRECTOR_02.jpg", 2],
        ["අතිරේක අධ්‍යක්ෂ", "Additional Director", "අතිරේක අධ්‍යක්ෂ (Add. Dir)", "රාජ්‍ය මූල්‍ය කළමනාකරණය පිළිබඳ පශ්චාත් උපාධිධාරියෙකි.", "DIRECTOR_03.jpg", 3],
        ["ප්‍රධාන ගණකාධිකාරී", "Chief Accountant", "මුදල් හා ආයෝජන (Finance)", "ලංකා බැංකුවේ සහ භාණ්ඩාගාරයේ වසර 15ක සේවා පළපුරුද්දක් ඇත.", "", 4],
        ["අධ්‍යක්ෂ (මෙහෙයුම්)", "Director (Ops)", "මෙහෙයුම් (Operations)", "ප්‍රාදේශීය ලේකම් කාර්යාල ජාලය සම්බන්ධීකරණය සහ මහජන සබඳතා පිළිබඳ විද්වතෙකි.", "", 5],
        ["අධ්‍යක්ෂ (තොරතුරු තාක්ෂණ)", "Director (IT)", "තොරතුරු තාක්ෂණ (IT)", "ඩිජිටල් පරිවර්තනය සහ දත්ත සුරක්ෂිතතාව පිළිබඳ පශ්චාත් උපාධිධාරී ඉංජිනේරුවරයෙකි.", "", 6]
    ])

    # 05. Events
    setup_sheet("DIV_05_EVENTS", VALIDATION_RULES["DIV_05_EVENTS"]["required_columns"], [
        ["2026-01-15", "ගාල්ල දිස්ත්‍රික් දැනුවත් කිරීමේ වැඩසටහන", "දකුණු පළාතේ ස්වයං රැකියා නියුක්තිකයින් සවිබල ගැන්වීම සහ සමාජ ආරක්ෂණ ජාලය පිළිබඳව දැනුවත් කිරීමේ මහා සම්මන්ත්‍රණය ගාල්ල ප්‍රධාන ශාලාවේදී අතිසාර්ථකව පැවැත්විණි.", "Galle", 450, "2026-01-15_GALLE"],
        ["2026-02-28", "කුරුණෑගල නව සාමාජිකයින් ලියාපදිංචිය", "වයඹ පළාතේ කෘෂිකාර්මික අංශයේ නියැලී සිටින ජනතාව ඉලක්ක කරගනිමින් සංවිධානය කල විශේෂ සාමාජික ප්‍රවර්ධන වැඩසටහන කුරුණෑගල දිස්ත්‍රික් ලේකම් කාර්යාලයේදී පැවැත්විණි.", "Kurunegala", 320, "2026-02-28_KURUNEGALA"],
        ["2026-02-05", "ඩිජිටල් තාක්ෂණ හඳුන්වාදීමේ වැඩසටහන", "දිවයින පුරා සිටින ප්‍රාදේශීය සම්බන්ධීකාරකවරුන් සඳහා නව මෘදුකාංගය පිළිබඳව පුහුණු වැඩසටහනක් කොළඹ ප්‍රධාන කාර්යාලයේදී පැවැත්විණි.", "Colombo", 150, ""],
        ["2026-02-14", "මාතර විශ්‍රාමිකයින්ගේ හමුව", "මාතර දිස්ත්‍රික්කයේ විශ්‍රාම වැටුප් ලබන ජ්‍යෙෂ්ඨ පුරවැසියන් සඳහා වාර්ෂික සුහද හමුව සහ ආගමික වැඩසටහන.", "Matara", 200, ""],
        ["2026-02-20", "මහනුවර ආයෝජන සමුළුව", "මධ්‍යම පළාතේ ව්‍යාපාරිකයින් සහ ස්වයං රැකියා නියුක්තිකයින් විශ්‍රාම අරමුදල් ආයෝජනයට යොමු කිරීමේ සම්මන්ත්‍රණය.", "Kandy", 500, ""]
    ])

    # 06. HR
    setup_sheet("DIV_06_HR", VALIDATION_RULES["DIV_06_HR"]["required_columns"], [
        ["2026-01", 350, 12, 2, 5, 120],
        ["2026-02", 360, 5, 1, 2, 80],
        ["2026-03", 364, 8, 4, 10, 240]
    ])

    # 07. Training
    setup_sheet("DIV_07_TRAINING", VALIDATION_RULES["DIV_07_TRAINING"]["required_columns"], [
        ["2026-01-10", "තොරතුරු තාක්ෂණ පුහුණුව (IT System Training)", 45, 2, "Colombo", 150000],
        ["2026-02-12", "නායකත්ව සංවර්ධන (Leadership Dev)", 30, 3, "Kandy", 200000],
        ["2026-02-18", "මූල්‍ය කළමනාකරණය (Financial Management)", 25, 1, "Galle", 85000],
        ["2026-02-25", "පාරිභෝගික සේවා (Customer Service)", 50, 2, "Kurunegala", 120000],
        ["2026-03-05", "කාර්යාල කළමනාකරණය (Office Administration)", 40, 2, "Colombo", 95000]
    ])

    # 08. Pensions
    setup_sheet("DIV_08_PENSIONS", VALIDATION_RULES["DIV_08_PENSIONS"]["required_columns"], [
        ["2026-01", 12500, 45000000, 2500000, 1500000],
        ["2026-02", 12650, 45500000, 2100000, 1200000],
        ["2026-03", 12820, 46100000, 2800000, 1800000]
    ])

    # 09. IT
    setup_sheet("DIV_09_IT", VALIDATION_RULES["DIV_09_IT"]["required_columns"], [
        ["Online Payment Gateway Integration", "In Progress", "2026-01-05", 85, 5000000, "Integrating BOC and Peoples Bank internet banking systems."],
        ["Document Management System", "Completed", "2025-10-01", 100, 3500000, "Digitizing physical files of pensioners."],
        ["Mobile App Development (Android/iOS)", "In Progress", "2026-02-01", 30, 8000000, "Building mobile application for member self-service."],
        ["Server Infrastructure Upgrade", "Completed", "2026-01-15", 100, 12000000, "Migrating internal databases to state-of-the-art cloud servers."],
        ["Cybersecurity Audit Rectification", "In Progress", "2026-02-20", 45, 1500000, "Implementing firewall rules and VPN for regional offices."]
    ])

    # 10. Audit
    setup_sheet("DIV_10_AUDIT", VALIDATION_RULES["DIV_10_AUDIT"]["required_columns"], [
        ["2026-01-20", "Finance", "Delay in bank reconciliation", "Medium", "Open", "Implementing automated reconciliation system."],
        ["2026-02-15", "IT", "Firewall policies not updated", "High", "Closed", "Updated policies as per 2026 security guidelines."],
        ["2026-02-18", "HR", "Incomplete personal files", "Low", "In Progress", "Instructed HR team to update missing NIC copies."],
        ["2026-02-25", "Procurement", "Tender documentation missing signatures", "High", "Open", "Called for an immediate inquiry board."],
        ["2026-02-28", "Regional Office (Galle)", "Petty cash book not maintained", "Medium", "Closed", "Issued stern warning and assigned a new accounting officer."]
    ])

    MASTER_EXCEL_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(MASTER_EXCEL_PATH)
    print(f"✅ Generated rich mock Excel file at: {MASTER_EXCEL_PATH}")

if __name__ == "__main__":
    create_mock_excel()
